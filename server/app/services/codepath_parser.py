r"""codepath.xlsx 파서 — 바코드 → 이미지 경로 매핑.

파일 형식: 헤더 없음, 2컬럼 고정
  A: 바코드 (EAN-13)
  B: Windows 절대경로 (Z:\물류부\scan\img\xxx.jpg)
"""

import json
import logging
import re
import time
from pathlib import PurePosixPath, PureWindowsPath

import openpyxl

logger = logging.getLogger(__name__)

_PATH_PREFIX_RE = re.compile(r"^.*?\\scan\\", re.IGNORECASE)

BATCH_SIZE = 500


def _to_relative_path(windows_path: str) -> str:
    """Windows 절대경로 → 상대경로 (img/xxx.jpg)."""
    cleaned = _PATH_PREFIX_RE.sub("", windows_path)
    return str(PurePosixPath(PureWindowsPath(cleaned)))


async def parse_codepath(db, file_path: str) -> dict:
    """codepath.xlsx를 파싱하여 BARCODE + IMAGE 테이블에 적재."""
    start = time.perf_counter()
    stats = {"added": 0, "updated": 0, "skipped": 0, "errors": 0, "error_details": []}

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    barcode_batch = []
    image_batch = []
    record_count = 0

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        try:
            barcode = str(row[0]).strip() if row[0] else None
            raw_path = str(row[1]).strip() if len(row) > 1 and row[1] else None

            if not barcode:
                stats["skipped"] += 1
                continue

            record_count += 1
            barcode_batch.append((barcode,))

            if raw_path and raw_path.lower() != "none":
                relative_path = _to_relative_path(raw_path)
                image_type = "real" if "real_image" in relative_path else "thumbnail"
                image_batch.append((barcode, relative_path, image_type))

            if len(barcode_batch) >= BATCH_SIZE:
                added, updated = await _flush_batch(db, barcode_batch, image_batch)
                stats["added"] += added
                stats["updated"] += updated
                barcode_batch.clear()
                image_batch.clear()

        except Exception as e:
            stats["errors"] += 1
            stats["error_details"].append(f"Row {row_idx}: {e}")
            if stats["errors"] <= 10:
                logger.warning("codepath row %d 파싱 실패: %s", row_idx, e)

    if barcode_batch:
        added, updated = await _flush_batch(db, barcode_batch, image_batch)
        stats["added"] += added
        stats["updated"] += updated

    wb.close()
    await db.commit()

    duration_ms = int((time.perf_counter() - start) * 1000)

    await db.execute(
        "INSERT INTO parse_log (file_name, file_type, record_count, added_count, "
        "updated_count, skipped_count, error_count, errors, duration_ms) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (
            str(file_path),
            "codepath",
            record_count,
            stats["added"],
            stats["updated"],
            stats["skipped"],
            stats["errors"],
            json.dumps(stats["error_details"][:50], ensure_ascii=False),
            duration_ms,
        ),
    )
    await db.commit()

    logger.info(
        "codepath 파싱 완료: 추가=%d, 갱신=%d, 스킵=%d, 에러=%d (%dms)",
        stats["added"], stats["updated"], stats["skipped"], stats["errors"], duration_ms,
    )
    return stats


async def _flush_batch(
    db, barcode_batch: list[tuple], image_batch: list[tuple]
) -> tuple[int, int]:
    barcodes = [b[0] for b in barcode_batch]
    placeholders = ",".join("?" * len(barcodes))
    cursor = await db.execute(
        f"SELECT COUNT(*) FROM barcode WHERE barcode IN ({placeholders})", barcodes
    )
    existing_count = (await cursor.fetchone())[0]
    added = len(barcodes) - existing_count
    updated = existing_count

    await db.executemany(
        "INSERT INTO barcode (barcode) VALUES (?) "
        "ON CONFLICT(barcode) DO UPDATE SET updated_at = datetime('now') "
        "WHERE sku_id IS NULL OR sku_id = ''",
        barcode_batch,
    )

    if image_batch:
        await db.executemany(
            "INSERT INTO image (barcode, file_path, image_type) VALUES (?, ?, ?) "
            "ON CONFLICT(barcode, file_path) DO UPDATE SET "
            "image_type = excluded.image_type, updated_at = datetime('now')",
            image_batch,
        )

    return added, updated
